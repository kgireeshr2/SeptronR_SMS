from datetime import datetime
from typing import List, Optional
from uuid import UUID
import csv
import io

from fastapi import HTTPException, UploadFile
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.auth import User
from app.repositories.student_repository import StudentRepository
from app.schemas.phase5 import (
    BulkImportResult,
    IssueTCRequest,
    PromoteStudentsRequest,
    StudentCreate,
    StudentDocumentCreate,
    StudentEnrollmentCreate,
    StudentParentCreate,
    StudentUpdate,
)
from app.utils.admission_number import generate_admission_number
from app.core.security import hash_password


class StudentService:
    def __init__(self, session: AsyncSession):
        self.session = session
        self.repository = StudentRepository(session)

    async def create_student(
        self, school_id: UUID, student_data: StudentCreate, current_user: User
    ):
        """Create a new student with auto-generated admission number if not provided."""
        # Generate admission number if not provided
        if not student_data.admission_number:
            admission_number = await generate_admission_number(self.session, school_id)
        else:
            # Validate uniqueness
            existing = await self.repository.get_by_admission_number(
                student_data.admission_number, school_id
            )
            if existing:
                raise HTTPException(
                    status_code=400,
                    detail=f"Admission number {student_data.admission_number} already exists",
                )
            admission_number = student_data.admission_number

        # Create student
        student = await self.repository.create(school_id, student_data, admission_number)

        # Auto-create student user account (username = admission number, password = DOB as DDMMYYYY)
        student_user_id = await self._create_student_login(school_id, student, admission_number, student_data.date_of_birth)
        if student_user_id:
            student.user_id = student_user_id
            await self.session.flush()

        # Create enrollment if provided
        if student_data.enrollment:
            await self.repository.add_enrollment(
                student.id, school_id, student_data.enrollment
            )

        # Create parents if provided
        for parent_data in student_data.parents:
            # Auto-create parent login whenever phone or email is available
            user_id = None
            if parent_data.phone or parent_data.email:
                user_id = await self._create_parent_login(school_id, parent_data, student.id)

            await self.repository.add_parent(student.id, school_id, parent_data, user_id)

        await self.session.commit()
        return await self.repository.get_detail(student.id, school_id)

    async def _create_student_login(
        self, school_id: UUID, student: any, admission_number: str, date_of_birth
    ) -> Optional[UUID]:
        """Auto-create a user account for a student (username = admission_number, password = DOB as DDMMYYYY)."""
        from sqlalchemy import select
        from app.models.auth import User
        from app.models.rbac import Role, UserRole as UserRoleModel
        from datetime import datetime, timezone

        # Skip if user already exists with this admission number as username
        result = await self.session.execute(
            select(User).where(User.username == admission_number, User.school_id == str(school_id))
        )
        if result.scalar_one_or_none():
            return None

        # Default password = DOB formatted as DDMMYYYY
        dob = date_of_birth
        default_password = dob.strftime("%d%m%Y") if hasattr(dob, "strftime") else str(dob).replace("-", "")

        user = User(
            school_id=school_id,
            username=admission_number,
            password_hash=hash_password(default_password),
            is_active=True,
        )
        self.session.add(user)
        await self.session.flush()

        # Assign "student" role if it exists
        role_result = await self.session.execute(
            select(Role).where(Role.school_id == str(school_id), Role.slug == "student")
        )
        student_role = role_result.scalar_one_or_none()
        if student_role:
            self.session.add(UserRoleModel(
                user_id=user.id,
                role_id=student_role.id,
                assigned_at=datetime.now(timezone.utc),
            ))
            await self.session.flush()

        return user.id

    async def _create_parent_login(
        self, school_id: UUID, parent_data: StudentParentCreate, student_id: UUID
    ) -> Optional[UUID]:
        """Create user account for parent portal access (auto-created when phone or email is present)."""
        from sqlalchemy import select
        from app.models.auth import User
        from app.models.rbac import Role, UserRole as UserRoleModel
        from datetime import datetime, timezone

        # Check for existing user by phone first, then email
        existing_user = None
        if parent_data.phone:
            r = await self.session.execute(
                select(User).where(User.phone == parent_data.phone, User.school_id == str(school_id))
            )
            existing_user = r.scalar_one_or_none()
        if not existing_user and parent_data.email:
            r = await self.session.execute(
                select(User).where(User.email == parent_data.email, User.school_id == str(school_id))
            )
            existing_user = r.scalar_one_or_none()
        if existing_user:
            return existing_user.id

        # Default credentials: username = phone (or email prefix), password = "Parent@123"
        default_password = "Parent@123"
        if parent_data.phone:
            import re
            digits = re.sub(r"\D", "", parent_data.phone)
            username = digits[-10:] if len(digits) >= 10 else digits
        elif parent_data.email:
            username = parent_data.email.split("@")[0]
        else:
            return None

        # Ensure username uniqueness within school
        r = await self.session.execute(
            select(User).where(User.username == username, User.school_id == str(school_id))
        )
        if r.scalar_one_or_none():
            suffix = str(student_id)[:8]
            username = f"{username}_{suffix}"

        user = User(
            school_id=school_id,
            username=username,
            phone=parent_data.phone or None,
            email=parent_data.email or None,
            password_hash=hash_password(default_password),
            is_active=True,
        )
        self.session.add(user)
        await self.session.flush()

        # Assign "parent" role
        role_result = await self.session.execute(
            select(Role).where(Role.school_id == str(school_id), Role.slug == "parent")
        )
        parent_role = role_result.scalar_one_or_none()
        if parent_role:
            self.session.add(UserRoleModel(
                user_id=user.id,
                role_id=parent_role.id,
                assigned_at=datetime.now(timezone.utc),
            ))
            await self.session.flush()

        return user.id

    async def update_student(
        self, student_id: UUID, school_id: UUID, student_data: StudentUpdate
    ):
        """Update student details."""
        student = await self.repository.update(student_id, school_id, student_data)
        if not student:
            raise HTTPException(status_code=404, detail="Student not found")

        await self.session.commit()
        return await self.repository.get_detail(student_id, school_id)

    async def delete_student(self, student_id: UUID, school_id: UUID, current_user: User):
        """Soft delete student."""
        success = await self.repository.soft_delete(student_id, school_id, current_user.id)
        if not success:
            raise HTTPException(status_code=404, detail="Student not found")

        await self.session.commit()
        return {"message": "Student deleted successfully"}

    async def promote_students(
        self, school_id: UUID, request: PromoteStudentsRequest, current_user: User
    ):
        """Bulk promote students to next academic year and class."""
        promotions = []

        for student_id in request.student_ids:
            # Get current enrollment
            current_enrollment = await self.repository.get_current_enrollment(
                student_id, school_id
            )
            if not current_enrollment:
                continue

            # Mark current enrollment as inactive
            await self.repository.mark_enrollment_inactive(
                current_enrollment.id, school_id
            )

            # Create new enrollment
            new_enrollment_data = StudentEnrollmentCreate(
                academic_year_id=request.to_academic_year_id,
                class_id=request.to_class_id,
                section_id=request.to_section_id,
                roll_number=None,  # Will be assigned later
                is_current=True,
            )
            await self.repository.add_enrollment(
                student_id, school_id, new_enrollment_data
            )

            # Record promotion
            promotion = await self.repository.promote_student(
                student_id=student_id,
                school_id=school_id,
                from_class_id=current_enrollment.class_id,
                to_class_id=request.to_class_id,
                from_year_id=request.from_academic_year_id,
                to_year_id=request.to_academic_year_id,
                promoted_by=current_user.id,
            )
            promotions.append(promotion)

        await self.session.commit()
        return promotions

    async def issue_transfer_certificate(
        self,
        student_id: UUID,
        school_id: UUID,
        tc_request: IssueTCRequest,
        current_user: User,
    ):
        """Issue transfer certificate and mark student as inactive."""
        student = await self.repository.get_by_id(student_id, school_id)
        if not student:
            raise HTTPException(status_code=404, detail="Student not found")

        # Mark student as inactive
        await self.repository.update(
            student_id, school_id, StudentUpdate(is_active=False)
        )

        # Issue TC
        transfer = await self.repository.issue_transfer_certificate(
            student_id=student_id,
            school_id=school_id,
            tc_number=tc_request.transfer_certificate_no,
            leaving_date=tc_request.leaving_date,
            reason=tc_request.reason,
            issued_by=current_user.id,
        )

        await self.session.commit()
        return transfer

    @staticmethod
    def _parse_import_rows(file: UploadFile, content: bytes) -> List[dict]:
        """Parse an uploaded CSV or XLSX file into a list of {header: value} dicts.

        XLSX is detected by filename extension or the ZIP magic bytes (PK\\x03\\x04);
        everything else is treated as UTF-8 CSV.
        """
        filename = (file.filename or "").lower()
        is_xlsx = filename.endswith(".xlsx") or content[:4] == b"PK\x03\x04"

        def _norm_header(h) -> str:
            # Drop the visual "required" marker so "first_name *" maps to "first_name".
            return ("" if h is None else str(h)).strip().rstrip("*").strip()

        if is_xlsx:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows_iter)
            except StopIteration:
                return []
            headers = [_norm_header(h) for h in header_row]

            def _cell(v):
                # openpyxl returns datetime objects for date-typed cells
                if isinstance(v, datetime):
                    return v.date().isoformat()
                if hasattr(v, "isoformat"):  # date
                    return v.isoformat()
                if v is None:
                    return ""
                return str(v).strip()

            rows = []
            for raw in rows_iter:
                if raw is None or all(c is None or str(c).strip() == "" for c in raw):
                    continue  # skip blank rows
                rows.append({headers[i]: _cell(raw[i]) for i in range(len(headers)) if headers[i]})
            return rows

        # CSV fallback
        csv_file = io.StringIO(content.decode("utf-8-sig"))
        return [
            {_norm_header(k): (v.strip() if isinstance(v, str) else v) for k, v in row.items()}
            for row in csv.DictReader(csv_file)
        ]

    async def bulk_import_students(
        self, school_id: UUID, file: UploadFile, current_user: User
    ) -> BulkImportResult:
        """Bulk import students (CSV or XLSX) including enrollment, parents and fees.

        Recognised columns (header row, case-sensitive):
          Student : first_name*, last_name*, date_of_birth*, gender*, blood_group,
                    religion, category, nationality, admission_date*, admission_number,
                    phone, email, address, city, state, pincode
          Class   : class_name, section_name, roll_number  (uses the current academic year)
          Parents : father_name/father_phone/father_email/father_occupation,
                    mother_name/mother_phone/mother_email/mother_occupation,
                    guardian_name/guardian_phone/guardian_email/guardian_relation
          Fees    : fee_master_name  (assigns the matching FeeMaster for the class)
        (* = required)
        """
        from app.models.students import ParentRelation
        from app.repositories.academic_repository import AcademicYearRepository
        from app.repositories.class_repository import ClassRepository, SectionRepository
        from app.repositories.fee_v2_repository import FeeV2Repository
        from app.schemas.fees_v2 import StudentFeeMasterAssignRequest

        content = await file.read()
        rows = self._parse_import_rows(file, content)

        # ----- Prefetch lookups (resolved to plain string ids so they survive commits) -----
        year_repo = AcademicYearRepository(self.session)
        class_repo = ClassRepository(self.session)
        section_repo = SectionRepository(self.session)
        fee_repo = FeeV2Repository(self.session)

        current_year = await year_repo.get_current(school_id)
        year_id = str(current_year.id) if current_year else None

        class_by_name: dict = {}
        if year_id:
            for cls in await class_repo.list_by_school_year(str(school_id), year_id):
                class_by_name[cls.name.strip().lower()] = str(cls.id)

        section_cache: dict = {}      # class_id -> {section_name_lower: section_id}
        fee_master_cache: dict = {}   # class_id -> [(name_lower, master_id)]

        async def resolve_sections(class_id: str) -> dict:
            if class_id not in section_cache:
                section_cache[class_id] = {
                    s.name.strip().lower(): str(s.id)
                    for s in await section_repo.list_by_class(class_id)
                }
            return section_cache[class_id]

        async def resolve_fee_masters(class_id: str) -> list:
            if class_id not in fee_master_cache:
                masters = await fee_repo.list_fee_masters(str(school_id), class_id=class_id, year_id=year_id)
                fee_master_cache[class_id] = [
                    (m["master"].name.strip().lower(), str(m["master"].id)) for m in masters
                ]
            return fee_master_cache[class_id]

        def _val(row: dict, key: str) -> Optional[str]:
            v = row.get(key)
            v = v.strip() if isinstance(v, str) else v
            return v or None

        def _build_parent(row, prefix, relation, is_primary):
            name = _val(row, f"{prefix}_name")
            if not name:
                return None
            rel_raw = _val(row, f"{prefix}_relation") or relation
            try:
                rel = ParentRelation(rel_raw.lower())
            except ValueError:
                rel = ParentRelation.other
            return StudentParentCreate(
                relation=rel,
                name=name,
                phone=_val(row, f"{prefix}_phone"),
                email=_val(row, f"{prefix}_email"),
                occupation=_val(row, f"{prefix}_occupation"),
                is_primary_contact=is_primary,
            )

        success_count = 0
        error_count = 0
        errors: list = []

        for row_num, row in enumerate(rows, start=2):  # row 1 is the header
            try:
                # ----- Enrollment (class / section) -----
                enrollment = None
                class_id = None
                class_name = _val(row, "class_name")
                if class_name:
                    if not year_id:
                        raise ValueError("No current academic year is set for this school")
                    class_id = class_by_name.get(class_name.lower())
                    if not class_id:
                        raise ValueError(f"Class '{class_name}' not found for the current academic year")

                    section_id = None
                    section_name = _val(row, "section_name")
                    if section_name:
                        sections = await resolve_sections(class_id)
                        section_id = sections.get(section_name.lower())
                        if not section_id:
                            raise ValueError(f"Section '{section_name}' not found in class '{class_name}'")

                    enrollment = StudentEnrollmentCreate(
                        academic_year_id=year_id,
                        class_id=class_id,
                        section_id=section_id,
                        roll_number=_val(row, "roll_number"),
                    )

                # ----- Fee structure (resolve BEFORE creating the student so a bad
                # fee name fails the row cleanly instead of orphaning a student) -----
                fee_master_id = None
                if class_id:
                    fee_master_name = _val(row, "fee_master_name")
                    masters = await resolve_fee_masters(class_id)
                    if fee_master_name:
                        fee_master_id = next(
                            (mid for name, mid in masters if name == fee_master_name.lower()), None
                        )
                        if not fee_master_id:
                            raise ValueError(
                                f"Fee structure '{fee_master_name}' not found for class '{class_name}'"
                            )
                    elif len(masters) == 1:
                        # No name given but the class has exactly one fee structure → auto-assign it.
                        fee_master_id = masters[0][1]

                # ----- Parents -----
                parents = []
                father = _build_parent(row, "father", "father", is_primary=True)
                if father:
                    parents.append(father)
                mother = _build_parent(row, "mother", "mother", is_primary=not father)
                if mother:
                    parents.append(mother)
                guardian = _build_parent(row, "guardian", "guardian", is_primary=not father and not mother)
                if guardian:
                    parents.append(guardian)

                # ----- Student -----
                student_data = StudentCreate(
                    first_name=row["first_name"],
                    last_name=row["last_name"],
                    date_of_birth=datetime.strptime(row["date_of_birth"], "%Y-%m-%d").date(),
                    gender=row["gender"],
                    blood_group=_val(row, "blood_group"),
                    religion=_val(row, "religion"),
                    category=_val(row, "category"),
                    nationality=_val(row, "nationality") or "Indian",
                    admission_date=datetime.strptime(row["admission_date"], "%Y-%m-%d").date(),
                    admission_number=_val(row, "admission_number"),
                    phone=_val(row, "phone"),
                    email=_val(row, "email"),
                    address=_val(row, "address"),
                    city=_val(row, "city"),
                    state=_val(row, "state"),
                    pincode=_val(row, "pincode"),
                    enrollment=enrollment,
                    parents=parents,
                )

                student = await self.create_student(school_id, student_data, current_user)
                student_id = str(student.id)

                # ----- Assign the resolved fee structure (student now exists) -----
                if fee_master_id:
                    await fee_repo.assign_fee_master(
                        str(school_id),
                        StudentFeeMasterAssignRequest(
                            student_ids=[student_id],
                            fee_master_id=fee_master_id,
                            academic_year_id=year_id,
                        ),
                        assigned_by=str(current_user.id),
                    )
                    await self.session.commit()

                success_count += 1

            except Exception as e:
                await self.session.rollback()
                error_count += 1
                errors.append({"row": row_num, "error": str(e)})

        await self.session.commit()

        return BulkImportResult(
            success_count=success_count,
            error_count=error_count,
            errors=errors,
        )

    async def export_students_csv(self, school_id: UUID) -> str:
        """Export students to CSV format."""
        students, _ = await self.repository.list(school_id, skip=0, limit=100000)

        output = io.StringIO()
        fieldnames = [
            "admission_number",
            "first_name",
            "last_name",
            "date_of_birth",
            "gender",
            "blood_group",
            "religion",
            "category",
            "nationality",
            "admission_date",
            "is_active",
        ]
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()

        for student in students:
            writer.writerow(
                {
                    "admission_number": student.admission_number,
                    "first_name": student.first_name,
                    "last_name": student.last_name,
                    "date_of_birth": student.date_of_birth.isoformat(),
                    "gender": student.gender,
                    "blood_group": student.blood_group or "",
                    "religion": student.religion or "",
                    "category": student.category or "",
                    "nationality": student.nationality,
                    "admission_date": student.admission_date.isoformat(),
                    "is_active": student.is_active,
                }
            )

        return output.getvalue()

